# Copyright 2026 Ihor Kozlitin
# SPDX-License-Identifier: Apache-2.0

import io
from pathlib import Path
from fastapi import FastAPI, HTTPException, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pypdf import PdfReader
from docx import Document

from app.schemas.api import (
    AnalyzeRequest,
    AnalyzeResponse,
    PseudonymizeRequest,
    PseudonymizeResponse,
    RestoreRequest,
    RestoreResponse,
    ExtractPDFResponse,
    ExtractWordResponse,
    DetectedEntity,
)
from app.core.analyzer import analyze_text
from app.core.anonymizer import process_pseudonymization, process_restoration

app = FastAPI(
    title="PIIUA - Ukrainian PII Pseudonymization & Restoration Service for Mediation/ODR",
    description="Privacy-preserving system for de-identifying personal data in Ukrainian legal and dispute texts using Presidio.",
    version="1.0.0",
)

# CORS Middleware
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"

if STATIC_DIR.exists():
    app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")


@app.get("/", include_in_schema=False)
async def serve_index():
    index_file = STATIC_DIR / "index.html"
    if index_file.exists():
        return FileResponse(str(index_file))
    return {"message": "PIIUA API is running. Access /docs for API documentation."}


@app.get("/health")
async def health_check():
    return {"status": "ok", "service": "PIIUA", "version": "1.0.0"}


@app.post("/api/v1/analyze", response_model=AnalyzeResponse)
async def analyze_pii(request: AnalyzeRequest):
    """
    Detect PII entities in Ukrainian text without modifying the text.
    """
    try:
        model_name = request.model_name or "uk_core_news_trf"
        raw_results = analyze_text(
            text=request.text,
            score_threshold=request.score_threshold,
            model_name=model_name
        )
        entities = [
            DetectedEntity(
                entity_type=r.entity_type,
                start=r.start,
                end=r.end,
                score=round(r.score, 2),
                text=request.text[r.start:r.end],
            )
            for r in raw_results
        ]
        return AnalyzeResponse(entities=entities)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/v1/pseudonymize", response_model=PseudonymizeResponse)
async def pseudonymize_pii(request: PseudonymizeRequest):
    """
    Pseudonymize Ukrainian text by replacing PII entities with structured reversible tokens (<PERSON_1>, <RNTRC_1>, etc.).
    Returns the pseudonymized text, the bidirectional token mapping dictionary, and detected entities.
    """
    try:
        model_name = request.model_name or "uk_core_news_trf"
        pseudo_text, mapping, entities = process_pseudonymization(
            text=request.text,
            score_threshold=request.score_threshold,
            model_name=model_name
        )
        return PseudonymizeResponse(
            pseudonymized_text=pseudo_text,
            mapping=mapping,
            entities=entities
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/v1/restore", response_model=RestoreResponse)
async def restore_pii(request: RestoreRequest):
    """
    Restore original PII values back into LLM responses or pseudonymized texts using the token mapping dictionary.
    """
    try:
        restored = process_restoration(
            pseudonymized_text=request.text,
            mapping=request.mapping
        )
        return RestoreResponse(restored_text=restored)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/v1/extract-pdf", response_model=ExtractPDFResponse)
async def extract_pdf_text(file: UploadFile = File(...)):
    """
    Extract text from an uploaded PDF document.
    """
    if not file.filename.lower().endswith('.pdf'):
        raise HTTPException(status_code=400, detail="Обраний файл не є PDF-документом.")
    
    try:
        content = await file.read()
        if not content:
            raise HTTPException(status_code=400, detail="Файл порожній.")
            
        reader = PdfReader(io.BytesIO(content))
        
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise HTTPException(status_code=400, detail="PDF-файл захищений паролем.")
        
        extracted_pages = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text and page_text.strip():
                extracted_pages.append(page_text.strip())
        
        full_text = "\n\n".join(extracted_pages)
        if not full_text:
            raise HTTPException(
                status_code=400,
                detail="Не вдалося витягти текст з PDF. Можливо, файл містить лише скановані зображення без текстового шару."
            )
        
        return ExtractPDFResponse(
            filename=file.filename,
            text=full_text,
            page_count=len(reader.pages)
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Помилка при зчитуванні PDF-файлу: {str(e)}")


@app.post("/api/v1/extract-word", response_model=ExtractWordResponse)
async def extract_word_text(file: UploadFile = File(...)):
    """
    Extract text from an uploaded Word document (.docx / .doc).
    """
    filename_lower = file.filename.lower()
    if not (filename_lower.endswith('.docx') or filename_lower.endswith('.doc')):
        raise HTTPException(status_code=400, detail="Обраний файл не є Word-документом (.docx, .doc).")
    
    try:
        content = await file.read()
        if not content:
            raise HTTPException(status_code=400, detail="Файл порожній.")
        
        full_text = ""
        if filename_lower.endswith('.docx'):
            try:
                doc = Document(io.BytesIO(content))
                paragraphs = []
                for p in doc.paragraphs:
                    if p.text and p.text.strip():
                        paragraphs.append(p.text.strip())
                
                if doc.tables:
                    for table in doc.tables:
                        for row in table.rows:
                            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                            if row_text:
                                paragraphs.append(" | ".join(row_text))
                
                full_text = "\n\n".join(paragraphs)
            except Exception as e:
                try:
                    import zipfile
                    import xml.etree.ElementTree as ET
                    with zipfile.ZipFile(io.BytesIO(content)) as z:
                        xml_content = z.read('word/document.xml')
                        tree = ET.fromstring(xml_content)
                        paragraphs = []
                        for p in tree.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
                            texts = [node.text for node in p.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t') if node.text]
                            if texts:
                                paragraphs.append("".join(texts).strip())
                        full_text = "\n\n".join([p for p in paragraphs if p])
                except Exception:
                    raise HTTPException(status_code=400, detail=f"Не вдалося зчитати DOCX файл: {str(e)}")
        else:
            try:
                import re
                raw_text = content.decode('utf-16le', errors='ignore')
                cleaned = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', raw_text)
                lines = [line.strip() for line in cleaned.splitlines() if len(line.strip()) > 3]
                if lines:
                    full_text = "\n\n".join(lines)
                else:
                    raise ValueError("Empty content")
            except Exception:
                raise HTTPException(
                    status_code=400,
                    detail="Файл у застарілому форматі .DOC не вдалося зчитати. Будь ласка, збережіть документ у форматі .DOCX та повторіть спробу."
                )

        if not full_text or not full_text.strip():
            raise HTTPException(
                status_code=400,
                detail="Не вдалося витягти текст з Word-документа. Можливо, файл порожній або захищений."
            )

        return ExtractWordResponse(
            filename=file.filename,
            text=full_text
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Помилка при зчитуванні Word-файлу: {str(e)}")


